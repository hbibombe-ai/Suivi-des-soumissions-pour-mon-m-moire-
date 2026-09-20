"""Tableaux de résultats du mémoire (chapitre IV), calculés en direct sur les données Kobo.

Chaque tableau reprend le numéro, le titre et la structure du chapitre « Résultats » :
tableaux descriptifs (effectif, %), tableaux bivariés (diarrhée oui/non, ORb [IC 95 %], p)
et modèle de régression logistique multivariée (ORa [IC 95 %], p, test de Hosmer-Lemeshow).
"""
from __future__ import annotations

VERSION_TDB = "5"  # doit correspondre à app.py
import io
import math
from typing import Callable, Dict, List, Optional, Tuple

import numpy as np
import pandas as pd
from scipy import stats

import viz

NA = "__na__"


# =========================================================== construction des variables
class Var:
    """Variable présentée dans un tableau : titre + série de modalités (ordonnées)."""
    def __init__(self, titre: str, serie: pd.Series, ordre: List[str], ref: Optional[str] = None,
                 multiple: bool = False, sous_pop: Optional[pd.Series] = None):
        self.titre, self.serie, self.ordre, self.ref = titre, serie, ordre, ref
        self.multiple, self.sous_pop = multiple, sous_pop


def col(df, c):
    return df[c] if c in df.columns else pd.Series(np.nan, index=df.index)


def choix(dico: dict, var: str) -> Dict[str, str]:
    return dico.get(var, {}).get("choix", {})


def v_code(df, dico, var, titre, codes=None, ref=None, labels=None) -> Var:
    """Variable à modalités codées dans le XLSForm (libellés et ordre du formulaire)."""
    lab = labels or choix(dico, var)
    codes = codes or list(lab.keys())
    s = col(df, var).map(lambda v: str(int(v)) if isinstance(v, (int, float)) and v == v and float(v).is_integer()
                         else (str(v) if v == v and v is not None else np.nan))
    s = s.map(lambda c: lab.get(c, c) if c == c else np.nan)
    ordre = [lab.get(str(c), str(c)) for c in codes]
    s = s.where(s.isin(ordre))
    return Var(titre, s, ordre, lab.get(str(ref)) if ref is not None else None)


def v_classes(df, serie, titre, bornes: List[Tuple[float, float, str]], ref=None) -> Var:
    def f(x):
        if x != x or x is None:
            return np.nan
        for lo, hi, nom in bornes:
            if lo <= x <= hi:
                return nom
        return np.nan
    return Var(titre, serie.astype(float).map(f), [b[2] for b in bornes], ref)


def v_groupes(df, serie, titre, groupes: List[Tuple[str, list]], ref=None) -> Var:
    """Regroupe des codes : [('Libellé', [codes…]), …]."""
    def f(x):
        if x != x or x is None:
            return np.nan
        for nom, codes in groupes:
            if x in codes:
                return nom
        return np.nan
    return Var(titre, serie.map(f), [g[0] for g in groupes], ref)


def v_bool(df, serie_bool, titre, oui="Oui", non="Non", ref=None, valide=None) -> Var:
    valide = serie_bool.notna() if valide is None else valide
    s = pd.Series(np.where(serie_bool.fillna(False).astype(bool), oui, non), index=serie_bool.index).where(valide)
    return Var(titre, s, [oui, non], ref)


def v_multiple(df, var, titre, dico) -> Var:
    """Question à réponses multiples (codes séparés par des espaces dans les exports Kobo)."""
    lab = choix(dico, var)
    return Var(titre, col(df, var), list(lab.items()), multiple=True)


# ============================================================================ calculs
def _lignes_desc(v: Var) -> List[dict]:
    rows = [{"Variables / modalités": v.titre, "Effectif (n)": "", "Pourcentage (%)": "", "_niv": "h"}]
    if v.multiple:
        s = v.serie.dropna().astype(str)
        n = len(s)
        for code, lab in v.ordre:
            k = sum(1 for x in s if code in str(x).split())
            rows.append({"Variables / modalités": lab, "Effectif (n)": k,
                         "Pourcentage (%)": viz.fr(100 * k / n) if n else "–", "_niv": "m"})
        rows[0]["Variables / modalités"] += f" (n = {n}, réponses multiples)"
        return rows
    s = v.serie.dropna()
    n = len(s)
    for m in v.ordre:
        k = int((s == m).sum())
        rows.append({"Variables / modalités": m, "Effectif (n)": k,
                     "Pourcentage (%)": viz.fr(100 * k / n) if n else "–", "_niv": "m"})
    rows[0]["Variables / modalités"] += f" (n = {n})"
    return rows


def tableau_descriptif(vars_: List[Var]) -> pd.DataFrame:
    rows = []
    for v in vars_:
        rows += _lignes_desc(v)
    return pd.DataFrame(rows)


def _or_ic(a, b, c, d):
    """OR (exposé a/b ; référence c/d) et IC 95 % de Woolf ; correction de Haldane si cellule nulle."""
    if min(a, b, c, d) == 0:
        a, b, c, d = a + .5, b + .5, c + .5, d + .5
    if min(a, b, c, d) <= 0:
        return (np.nan,) * 3
    o = (a * d) / (b * c)
    se = math.sqrt(1 / a + 1 / b + 1 / c + 1 / d)
    return o, o * math.exp(-1.96 * se), o * math.exp(1.96 * se)


def _p_valeur(tab: pd.DataFrame) -> Tuple[float, str]:
    tab = tab.loc[(tab.sum(axis=1) > 0), (tab.sum(axis=0) > 0)]
    if tab.shape[0] < 2 or tab.shape[1] < 2:
        return np.nan, ""
    chi2, p, _, att = stats.chi2_contingency(tab.values, correction=False)
    if tab.shape == (2, 2) and (att < 5).any():
        return stats.fisher_exact(tab.values)[1], "Fisher"
    return p, "Chi²"


def fmt_p(p):
    if p != p:
        return "–"
    return "< 0,001" if p < 0.001 else viz.fr(p, 3)


def tableau_bivarie(df: pd.DataFrame, vars_: List[Var], issue: str = "j1") -> Tuple[pd.DataFrame, list]:
    y = col(df, issue)
    rows, pvals = [], []
    for v in vars_:
        ok = v.serie.notna() & y.isin([0, 1])
        s, yy = v.serie[ok], y[ok]
        tab = pd.crosstab(s, yy).reindex(index=v.ordre, columns=[1, 0], fill_value=0)
        p, test = _p_valeur(tab)
        pvals.append((v.titre, p))
        rows.append({"Variables / modalités": v.titre, "Diarrhée : Oui n (%)": "", "Diarrhée : Non n (%)": "",
                     "ORb [IC 95 %]": "", "p": fmt_p(p) + (f" ({test})" if test == "Fisher" else ""), "_niv": "h"})
        ref = v.ref or v.ordre[0]
        c_, d_ = tab.loc[ref, 1], tab.loc[ref, 0]
        for m in v.ordre:
            a_, b_ = tab.loc[m, 1], tab.loc[m, 0]
            tot = a_ + b_
            if m == ref:
                orr = "1 (réf.)"
            else:
                o, lo, hi = _or_ic(a_, b_, c_, d_)
                orr = f"{viz.fr(o, 2)} [{viz.fr(lo, 2)} – {viz.fr(hi, 2)}]" if o == o else "–"
            rows.append({"Variables / modalités": m,
                         "Diarrhée : Oui n (%)": f"{a_} ({viz.fr(100 * a_ / tot)})" if tot else "0",
                         "Diarrhée : Non n (%)": f"{b_} ({viz.fr(100 * b_ / tot)})" if tot else "0",
                         "ORb [IC 95 %]": orr, "p": "", "_niv": "m"})
    return pd.DataFrame(rows), pvals


def hosmer_lemeshow(y, p, g=10):
    d = pd.DataFrame({"y": y, "p": p}).sort_values("p")
    d["grp"] = pd.qcut(d["p"].rank(method="first"), g, labels=False, duplicates="drop")
    obs = d.groupby("grp")["y"].sum(); att = d.groupby("grp")["p"].sum(); n = d.groupby("grp").size()
    chi = (((obs - att) ** 2) / (att * (1 - att / n))).sum()
    ddl = max(1, d["grp"].nunique() - 2)
    return chi, ddl, 1 - stats.chi2.cdf(chi, ddl)


def regression_logistique(df: pd.DataFrame, expositions: Dict[str, pd.Series], issue="j1"):
    """Régression logistique sur des expositions binaires (1 = exposé). Renvoie (tableau, infos)."""
    import statsmodels.api as sm
    X = pd.DataFrame(expositions).astype(float)
    y = col(df, issue)
    ok = X.notna().all(axis=1) & y.isin([0, 1])
    X, y = X[ok], y[ok].astype(int)
    X = X.loc[:, X.std() > 0]
    if len(y) < 30 or X.shape[1] == 0 or y.sum() < 5:
        return pd.DataFrame(), {"erreur": "Effectifs insuffisants pour ajuster le modèle."}
    Xc = sm.add_constant(X, has_constant="add")
    try:
        m = sm.Logit(y, Xc).fit(disp=0, maxiter=200)
    except Exception as e:  # noqa: BLE001 — séparation parfaite, colinéarité…
        return pd.DataFrame(), {"erreur": f"Le modèle n'a pas convergé ({e})."}
    ci = m.conf_int()
    rows = []
    for v in X.columns:
        rows.append({"Variables (exposé vs non exposé)": v, "ORa": viz.fr(math.exp(m.params[v]), 2),
                     "IC à 95 %": f"{viz.fr(math.exp(ci.loc[v, 0]), 2)} – {viz.fr(math.exp(ci.loc[v, 1]), 2)}",
                     "p": fmt_p(m.pvalues[v]), "_ora": math.exp(m.params[v]),
                     "_lo": math.exp(ci.loc[v, 0]), "_hi": math.exp(ci.loc[v, 1])})
    chi, ddl, p_hl = hosmer_lemeshow(y.values, m.predict(Xc).values)
    return pd.DataFrame(rows), {"n": int(len(y)), "cas": int(y.sum()), "hl": (chi, ddl, p_hl),
                                "pseudo_r2": m.prsquared}


# ======================================================================= variables
def variables(df: pd.DataFrame, dico: dict) -> Dict[str, object]:
    """Toutes les variables des tableaux du mémoire, calculées une seule fois."""
    V = {}
    b3 = col(df, "b3").astype(float)
    cas = col(df, "j1") == 1
    classes_age = [(0, 5, "0 à 5 mois"), (6, 11, "6 à 11 mois"), (12, 23, "12 à 23 mois"), (24, 59, "24 à 59 mois")]
    V["age"] = v_classes(df, b3, "Âge de l’enfant", classes_age, ref="24 à 59 mois")
    V["sexe"] = v_code(df, dico, "b1", "Sexe", ref=2)
    V["rang"] = v_classes(df, col(df, "b5"), "Rang de naissance", [(1, 1, "1"), (2, 3, "2 à 3"), (4, 99, "4 et plus")], ref="1")
    V["source_age"] = v_code(df, dico, "b4", "Source de vérification de l’âge")
    V["carnet"] = v_code(df, dico, "b6", "Carnet de vaccination disponible")
    V["rota"] = v_code(df, dico, "b7", "Vaccination contre le rotavirus", ref=1)
    # répondant
    V["lien"] = v_code(df, dico, "c1", "Lien avec l’enfant")
    V["age_rep"] = v_classes(df, col(df, "c2"), "Âge du répondant (années)",
                             [(0, 19, "Moins de 20"), (20, 29, "20 à 29"), (30, 39, "30 à 39"), (40, 120, "40 et plus")], ref="30 à 39")
    V["etudes"] = v_code(df, dico, "c3", "Niveau d’études", codes=["3", "2", "1", "0"], ref=3)
    V["activite"] = v_code(df, dico, "c4", "Activité principale")
    V["matrimonial"] = v_code(df, dico, "c5", "Situation matrimoniale")
    # ménage
    V["taille"] = v_classes(df, col(df, "a8"), "Taille du ménage (personnes)", [(1, 5, "1 à 5"), (6, 8, "6 à 8"), (9, 99, "9 et plus")], ref="1 à 5")
    V["enfants5"] = v_classes(df, col(df, "a9"), "Nombre d’enfants de 0 à 59 mois", [(1, 1, "1"), (2, 2, "2"), (3, 99, "3 et plus")])
    sb = col(df, "score_biens").astype(float)
    try:
        q = pd.qcut(sb.rank(method="first"), 5, labels=[f"Quintile {i}" for i in range(1, 6)])
    except ValueError:
        q = pd.Series(np.nan, index=df.index)
    lab_q = {"Quintile 1": "Quintile 1 (le plus pauvre)", "Quintile 5": "Quintile 5 (le plus riche)"}
    q = q.astype(object).map(lambda x: lab_q.get(x, x) if x == x else np.nan)
    V["quintile"] = Var("Niveau socio-économique (score des biens)", q,
                        ["Quintile 5 (le plus riche)", "Quintile 4", "Quintile 3", "Quintile 2", "Quintile 1 (le plus pauvre)"],
                        ref="Quintile 5 (le plus riche)")
    V["occupation"] = v_code(df, dico, "d14", "Statut d’occupation du logement")
    V["sol"] = v_code(df, dico, "d15", "Matériau du sol")
    V["sol2"] = v_groupes(df, col(df, "d15"), "Matériau du sol", [("Ciment/carrelage", [2, 3]), ("Terre/sable", [1])], ref="Ciment/carrelage")
    prom = col(df, "promiscuite").astype(float)
    if prom.isna().all():
        prom = col(df, "d17") / col(df, "d16")
    V["promiscuite"] = v_classes(df, prom, "Personnes par pièce à dormir", [(0, 3, "3 ou moins"), (3.0001, 99, "Plus de 3 (surpeuplement)")], ref="3 ou moins")
    # environnement
    V["inondation"] = v_code(df, dico, "d18", "Inondation au cours des 12 derniers mois", codes=["0", "1"], ref=0)
    V["freq_inond"] = v_code(df, dico, "d19", "Fréquence des inondations (si oui)")
    V["lieu_inond"] = v_code(df, dico, "d20", "Lieu atteint par l’eau (si oui)")
    V["stagnantes"] = v_code(df, dico, "d21", "Eaux stagnantes à proximité du logement", codes=["0", "1"], ref=0)
    V["caniveau"] = v_code(df, dico, "d22", "Caniveau ou système d’évacuation", codes=["1", "2", "0"], ref=1)
    V["depotoir"] = v_code(df, dico, "d23", "Distance du dépotoir le plus proche")
    V["depotoir2"] = v_groupes(df, col(df, "d23"), "Distance du dépotoir", [("Plus de 50 m ou aucun", [3, 4]), ("10 à 50 m", [2]), ("Moins de 10 m", [1])], ref="Plus de 50 m ou aucun")
    # alimentation
    V["deja_allaite"] = v_code(df, dico, "e1", "Déjà allaité au sein")
    V["allaite"] = v_code(df, dico, "e2", "Allaité actuellement", ref=1)
    V["exclusif"] = v_code(df, dico, "e3", "Allaitement exclusif pendant les six premiers mois", ref=1)
    intro = pd.Series(np.where(col(df, "e4_intro") == 0, "Pas encore introduits",
                               np.where(col(df, "e4") < 6, "Avant 6 mois", np.where(col(df, "e4") >= 6, "6 mois et plus", ""))),
                      index=df.index).replace("", np.nan)
    V["intro"] = Var("Âge d’introduction d’autres aliments ou liquides", intro, ["6 mois et plus", "Avant 6 mois", "Pas encore introduits"], ref="6 mois et plus")
    V["hors_menage"] = v_code(df, dico, "e5", "Aliments préparés hors du ménage", codes=["0", "2", "1"], ref=0)
    V["conservation"] = v_code(df, dico, "e7", "Mode de conservation des aliments (si conservés)")
    V["conservation2"] = v_groupes(df, col(df, "e7"), "Conservation des aliments", [("Récipient couvert / réfrigérateur", [1, 2]), ("Récipient non couvert", [3])], ref="Récipient couvert / réfrigérateur")
    # nutrition
    if "cl_whz" in df.columns:
        ordre_z = ["Normal (≥ −2)", "Modéré (−3 à < −2)", "Sévère (< −3)"]
        V["whz"] = Var("Émaciation (poids-pour-taille)", df["cl_whz"].where(df["cl_whz"].isin(ordre_z)), ordre_z)
        V["haz"] = Var("Retard de croissance (taille-pour-âge)", df["cl_haz"].where(df["cl_haz"].isin(ordre_z)), ordre_z)
        V["waz"] = Var("Insuffisance pondérale (poids-pour-âge)", df["cl_waz"].where(df["cl_waz"].isin(ordre_z)), ordre_z)
        ordre_pb = ["Normal (≥ 125 mm)", "MAM (115–124 mm)", "MAS (< 115 mm)", "MAS (œdèmes)"]
        V["pb"] = Var("Périmètre brachial / œdèmes (6 à 59 mois)", df["cl_pb"].where(df["cl_pb"].isin(ordre_pb) & ((b3 >= 6) | (df["cl_pb"] == "MAS (œdèmes)"))), ordre_pb)
        for k, t in [("emaciation", "Émaciation (P/T < −2 z)"), ("retard_croissance", "Retard de croissance (T/A < −2 z)"),
                     ("insuffisance_ponderale", "Insuffisance pondérale (P/A < −2 z)"), ("malnutrition_aigue_pb", "Malnutrition aiguë selon le PB (< 125 mm ou œdèmes)")]:
            V["b_" + k] = v_bool(df, df[k] == 1, t, ref="Non", valide=df[k].notna())
            V["b_" + k].ordre = ["Non", "Oui"]
    V["oedemes"] = v_code(df, dico, "m11", "Œdèmes bilatéraux")
    # eau
    V["source_eau"] = v_code(df, dico, "f1", "Source principale d’eau de boisson")
    ea = col(df, "eau_amelioree")
    V["eau_am"] = v_bool(df, ea == 1, "Source d’eau de boisson", oui="Améliorée", non="Non améliorée", ref="Améliorée", valide=ea.notna())
    V["jmp_eau"] = Var("Niveau de service JMP (eau)", col(df, "jmp_eau"), ["Basique", "Limité", "Non amélioré"], ref="Basique")
    V["temps_eau"] = v_code(df, dico, "f8", "Temps de collecte (aller-retour)")
    V["interruptions"] = v_code(df, dico, "f2", "Interruptions fréquentes d’approvisionnement", codes=["0", "1"], ref=0)
    V["traitement"] = v_code(df, dico, "f3", "Traitement de l’eau de boisson", codes=["1", "2", "0"], ref=1)
    V["methode"] = v_code(df, dico, "f4", "Méthode de traitement (si traitement)")
    V["recipient"] = v_code(df, dico, "f5", "Récipient de stockage")
    V["couvert"] = v_code(df, dico, "f6", "Récipient habituellement couvert", ref=1)
    V["prelevement"] = v_code(df, dico, "f7", "Mode de prélèvement de l’eau")
    V["prelevement2"] = v_groupes(df, col(df, "f7"), "Mode de prélèvement", [("Robinet / versée directement", [1, 2]), ("Gobelet / louche", [3]), ("Main ou récipient plongé", [4])], ref="Robinet / versée directement")
    # assainissement / hygiène / déchets
    V["latrine"] = v_code(df, dico, "g1", "Type de toilettes")
    V["latrine2"] = v_groupes(df, col(df, "g1"), "Type de toilettes", [("Amélioré", [1, 2, 3]), ("Non amélioré", [4, 6]), ("Défécation à l’air libre", [5])], ref="Amélioré")
    V["partage"] = v_code(df, dico, "g2", "Toilettes partagées avec d’autres ménages", codes=["0", "1"], ref=0)
    V["jmp_ass"] = Var("Niveau de service JMP (assainissement)", col(df, "jmp_assainissement"), ["Basique", "Limité", "Non amélioré", "Défécation en plein air"], ref="Basique")
    V["selles"] = v_code(df, dico, "g4", "Élimination des selles de l’enfant")
    V["selles2"] = v_groupes(df, col(df, "g4"), "Élimination des selles de l’enfant", [("Sûre (toilettes, enfouies)", [1, 2]), ("Non sûre", [3, 4, 5, 6])], ref="Sûre (toilettes, enfouies)")
    V["feces_obs"] = v_code(df, dico, "g5", "Matières fécales observées (latrine / cour)", codes=["0", "1"], ref=0)
    V["feces_obs_d"] = v_code(df, dico, "g5", "Matières fécales observées autour de la latrine ou dans la cour")
    V["ruissellement"] = v_code(df, dico, "g6", "Ruissellement atteignant la latrine lors des pluies", codes=["0", "1"], ref=0)
    V["dlm_toilettes"] = v_code(df, dico, "g7", "Dispositif de lavage des mains près des toilettes")
    V["lieu_lavage"] = v_code(df, dico, "i1", "Endroit habituel de lavage des mains")
    V["eau_obs"] = v_code(df, dico, "i2", "Eau disponible (observée)")
    V["savon"] = v_code(df, dico, "i3", "Savon ou produit nettoyant disponible")
    V["jmp_hyg"] = Var("Niveau de service JMP (hygiène)", col(df, "jmp_hygiene"), ["Basique", "Limité", "Aucun service"], ref="Basique")
    i4 = col(df, "i4").astype(str).where(col(df, "i4").notna())
    V["moments"] = v_multiple(df, "i4", "Moments de lavage des mains cités", dico)
    V["lav_toilettes"] = v_bool(df, i4.str.contains("apres_toilettes"), "Lavage des mains après les toilettes (cité)", ref="Oui", valide=i4.notna())
    V["lav_nourrir"] = v_bool(df, i4.str.contains("avant_nourrir"), "Lavage des mains avant de nourrir l’enfant (cité)", ref="Oui", valide=i4.notna())
    V["dechets"] = v_code(df, dico, "h1", "Mode d’élimination des déchets")
    V["dechets2"] = v_groupes(df, col(df, "h1"), "Mode d’élimination des déchets", [("Service organisé / dépotoir officiel", [1, 2]), ("Dépotoir sauvage", [3]), ("Brûlage / enfouissement", [4, 5]), ("Caniveau / rivière", [6])], ref="Service organisé / dépotoir officiel")
    V["dechets_fermes"] = v_code(df, dico, "h2", "Déchets stockés dans un récipient fermé", ref=1)
    # épisode (cas seulement)
    V["duree"] = v_classes(df, col(df, "j3").where(cas), "Durée de l’épisode", [(0, 3, "1 à 3 jours"), (4, 6, "4 à 6 jours"), (7, 13, "7 à 13 jours"), (14, 999, "14 jours et plus (persistante)")])
    V["sang"] = v_code(df, dico, "j4", "Sang visible dans les selles")
    V["fievre"] = v_code(df, dico, "j5", "Fièvre pendant l’épisode")
    V["vomi"] = v_code(df, dico, "j6", "Vomissements pendant l’épisode")
    # connaissances
    for k, t in [("l1", "A déjà entendu parler des SRO"), ("l2", "Rôle des SRO (si oui)"), ("l3", "Sait préparer un sachet de SRO"),
                 ("l4", "A déjà entendu parler du zinc"), ("l5", "Le zinc doit être administré s’il est recommandé"),
                 ("l6", "Continuer l’allaitement pendant la diarrhée"), ("l7", "Continuer à proposer des aliments appropriés"),
                 ("l8", "Proposer davantage de liquides")]:
        V[k] = v_code(df, dico, k, t)
    V["signes"] = v_multiple(df, "l9", "Signes de danger cités", dico)
    sc = col(df, "score_connaissances").astype(float)
    V["niveau_conn"] = v_classes(df, sc, "Niveau de connaissances (score sur 8)", [(6, 8, "Bon (6 à 8 points)"), (4, 5, "Moyen (4 à 5 points)"), (0, 3, "Faible (0 à 3 points)")])
    # prise en charge (cas)
    for k, t in [("k1", "Liquides reçus par rapport à l’habitude"), ("k2", "Allaitement pendant l’épisode"), ("k3", "Alimentation pendant l’épisode"),
                 ("k4", "SRO reçus"), ("k5", "Zinc reçu"), ("k7", "Soins ou conseils recherchés hors du domicile"),
                 ("k8", "Premier lieu de recours (si oui)"), ("k9", "Délai de recours (si oui)")]:
        V[k] = v_code(df, dico, k, t)
        V[k].serie = V[k].serie.where(cas)
    V["duree_zinc"] = v_classes(df, col(df, "k6"), "Durée du zinc (si reçu)", [(0, 9, "Moins de 10 jours"), (10, 14, "10 à 14 jours")])
    V["sro_zinc"] = v_bool(df, (col(df, "k4") == 1) & (col(df, "k5") == 1), "SRO et zinc associés", valide=cas)
    V["pec_adequate"] = v_bool(df, col(df, "prise_charge_adequate") == 1, "Prise en charge adéquate", valide=cas & col(df, "prise_charge_adequate").notna())
    V["raisons"] = v_multiple(df, "k10", "Raisons de l’absence de recours (réponses multiples)", dico)
    V["_score"] = sc
    return V


# ============================================================================ catalogue
NOTE_BIV = ("% en ligne. ORb : odds ratio brut ; IC 95 % : intervalle de confiance à 95 % (Woolf) ; réf. : modalité de référence ; "
            "p : test du Chi-deux de Pearson (test exact de Fisher si effectif théorique < 5).")

def catalogue(V: dict) -> List[dict]:
    """Liste ordonnée des tableaux du chapitre IV (numéro, section, titre, type, variables)."""
    g = lambda *k: [V[x] for x in k if x in V]
    return [
        {"num": 8, "section": "Prévalence", "titre": "Caractéristiques cliniques des épisodes diarrhéiques", "type": "desc", "vars": g("duree", "sang", "fievre", "vomi")},
        {"num": 9, "section": "Caractéristiques", "titre": "Caractéristiques des enfants enquêtés", "type": "desc", "vars": g("age", "sexe", "rang", "source_age", "carnet", "rota")},
        {"num": 10, "section": "Caractéristiques", "titre": "Caractéristiques des mères ou responsables des enfants", "type": "desc", "vars": g("lien", "age_rep", "etudes", "activite", "matrimonial")},
        {"num": 11, "section": "Caractéristiques", "titre": "Caractéristiques socio-économiques et d’habitat des ménages", "type": "desc", "vars": g("taille", "enfants5", "quintile", "occupation", "sol", "promiscuite")},
        {"num": 12, "section": "Caractéristiques", "titre": "Exposition des ménages aux inondations et à l’insalubrité de l’environnement", "type": "desc", "vars": g("inondation", "freq_inond", "lieu_inond", "stagnantes", "caniveau", "depotoir")},
        {"num": 13, "section": "Caractéristiques", "titre": "Pratiques d’allaitement et d’alimentation des enfants", "type": "desc", "vars": g("deja_allaite", "allaite", "exclusif", "intro", "hors_menage", "conservation")},
        {"num": 14, "section": "Caractéristiques", "titre": "État nutritionnel des enfants selon les normes de croissance de l’OMS", "type": "desc", "vars": g("whz", "haz", "waz", "pb", "oedemes"),
         "note": "z-scores calculés selon les normes OMS 2006 (méthode LMS) ; valeurs biologiquement invraisemblables exclues. Chiffres définitifs à confirmer avec WHO Anthro."},
        {"num": 15, "section": "Facteurs associés", "titre": "Diarrhée selon les caractéristiques de l’enfant", "type": "biv", "vars": g("age", "sexe", "rang", "rota")},
        {"num": 16, "section": "Facteurs associés", "titre": "Diarrhée selon les caractéristiques de la mère ou du responsable et du ménage", "type": "biv", "vars": g("etudes", "age_rep", "taille", "quintile", "promiscuite", "sol2")},
        {"num": 17, "section": "Facteurs associés", "titre": "Diarrhée selon l’allaitement et l’alimentation de l’enfant", "type": "biv", "vars": g("allaite", "exclusif", "intro", "hors_menage", "conservation2")},
        {"num": 18, "section": "Facteurs associés", "titre": "Diarrhée selon l’état nutritionnel de l’enfant", "type": "biv", "vars": g("b_emaciation", "b_retard_croissance", "b_insuffisance_ponderale", "b_malnutrition_aigue_pb"),
         "note": "Étude transversale : le sens de la relation entre diarrhée et état nutritionnel ne peut être établi."},
        {"num": 19, "section": "Facteurs associés", "titre": "Diarrhée selon l’approvisionnement en eau de boisson", "type": "biv", "vars": g("eau_am", "jmp_eau", "interruptions", "traitement", "couvert", "prelevement2")},
        {"num": 20, "section": "Facteurs associés", "titre": "Diarrhée selon l’assainissement, l’hygiène des mains et la gestion des selles de l’enfant", "type": "biv", "vars": g("latrine2", "partage", "selles2", "feces_obs", "ruissellement", "jmp_hyg", "lav_toilettes", "lav_nourrir")},
        {"num": 21, "section": "Facteurs associés", "titre": "Diarrhée selon la gestion des déchets et l’exposition aux inondations", "type": "biv", "vars": g("dechets2", "dechets_fermes", "inondation", "stagnantes", "caniveau", "depotoir2")},
        {"num": 23, "section": "Prévention et connaissances", "titre": "Accès à l’eau de boisson, traitement et stockage", "type": "desc", "vars": g("source_eau", "jmp_eau", "temps_eau", "interruptions", "traitement", "methode", "recipient", "couvert", "prelevement")},
        {"num": 24, "section": "Prévention et connaissances", "titre": "Assainissement et élimination des selles de l’enfant", "type": "desc", "vars": g("latrine", "partage", "jmp_ass", "selles", "feces_obs_d", "dlm_toilettes")},
        {"num": 25, "section": "Prévention et connaissances", "titre": "Hygiène des mains dans les ménages", "type": "desc", "vars": g("lieu_lavage", "eau_obs", "savon", "jmp_hyg", "moments")},
        {"num": 26, "section": "Prévention et connaissances", "titre": "Connaissances des répondants sur les SRO et le zinc", "type": "desc", "vars": g("l1", "l2", "l3", "l4", "l5")},
        {"num": 27, "section": "Prévention et connaissances", "titre": "Connaissances sur l’allaitement, l’alimentation et les liquides pendant la diarrhée", "type": "desc", "vars": g("l6", "l7", "l8")},
        {"num": 28, "section": "Prévention et connaissances", "titre": "Signes de danger cités par les répondants", "type": "desc", "vars": g("signes")},
        {"num": 29, "section": "Prévention et connaissances", "titre": "Niveau global de connaissances des répondants (score sur 8)", "type": "desc", "vars": g("niveau_conn")},
        {"num": 30, "section": "Prise en charge", "titre": "Prise en charge à domicile des enfants ayant eu la diarrhée", "type": "desc", "vars": g("k1", "k2", "k3", "k4", "k5", "duree_zinc", "sro_zinc", "pec_adequate")},
        {"num": 31, "section": "Prise en charge", "titre": "Recours aux soins pour l’épisode diarrhéique", "type": "desc", "vars": g("k7", "k8", "k9", "raisons")},
    ]


# Expositions binaires proposées pour le modèle multivarié (1 = exposé)
def expositions_binaires(df: pd.DataFrame) -> Dict[str, pd.Series]:
    b3 = col(df, "b3").astype(float)
    def bin_(cond, valide):
        return cond.astype(float).where(valide)
    E = {
        "Âge 6 à 23 mois": bin_((b3 >= 6) & (b3 <= 23), b3.notna()),
        "Mère sans instruction ou primaire": bin_(col(df, "c3").isin([0, 1]), col(df, "c3").notna()),
        "Surpeuplement (> 3 pers./pièce)": bin_((col(df, "d17") / col(df, "d16")) > 3, col(df, "d16").notna() & col(df, "d17").notna()),
        "Non vacciné contre le rotavirus": bin_(col(df, "b7") == 0, col(df, "b7").isin([0, 1])),
        "Pas d’allaitement exclusif (6 mois)": bin_(col(df, "e3") == 0, col(df, "e3").isin([0, 1])),
        "Source d’eau non améliorée": bin_(col(df, "eau_amelioree") == 0, col(df, "eau_amelioree").notna()),
        "Eau de boisson jamais traitée": bin_(col(df, "f3") == 0, col(df, "f3").notna()),
        "Récipient d’eau non couvert": bin_(col(df, "f6") == 0, col(df, "f6").notna()),
        "Toilettes non améliorées": bin_(col(df, "g1").isin([4, 5, 6]), col(df, "g1").notna()),
        "Toilettes partagées": bin_(col(df, "g2") == 1, col(df, "g2").notna()),
        "Élimination non sûre des selles de l’enfant": bin_(col(df, "g4").isin([3, 4, 5, 6]), col(df, "g4").notna()),
        "Pas de savon au point de lavage": bin_(col(df, "i3") != 1, col(df, "i1").notna()),
        "Inondation (12 derniers mois)": bin_(col(df, "d18") == 1, col(df, "d18").notna()),
        "Eaux stagnantes à proximité": bin_(col(df, "d21") == 1, col(df, "d21").notna()),
        "Dépotoir à moins de 10 m": bin_(col(df, "d23") == 1, col(df, "d23").notna()),
    }
    if "emaciation" in df.columns:
        E["Émaciation (P/T < −2 z)"] = df["emaciation"]
        E["Retard de croissance (T/A < −2 z)"] = df["retard_croissance"]
    return E


def p_bivarie(df: pd.DataFrame, s: pd.Series, issue="j1") -> float:
    y = col(df, issue)
    ok = s.notna() & y.isin([0, 1])
    return _p_valeur(pd.crosstab(s[ok], y[ok]))[0] if ok.sum() else np.nan


# ============================================================================ export
def exporter_excel(tableaux: List[Tuple[str, pd.DataFrame, str]]) -> bytes:
    """Un classeur : une feuille par tableau (titre, tableau, note), prêt à copier dans le mémoire."""
    from openpyxl.styles import Alignment, Font, PatternFill
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        for titre, t, note in tableaux:
            nom = titre.split(" :")[0].replace("Tableau ", "T")[:31]
            vis = t[[c for c in t.columns if not c.startswith("_")]]
            vis.to_excel(w, sheet_name=nom, index=False, startrow=2)
            ws = w.sheets[nom]
            ws["A1"] = titre; ws["A1"].font = Font(bold=True, size=12)
            for c in ws[3]:
                c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="D9E2F3")
                c.alignment = Alignment(wrap_text=True, vertical="center")
            if "_niv" in t.columns:
                for i, niv in enumerate(t["_niv"], start=4):
                    if niv == "h":
                        for c in ws[i]:
                            c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="F2F2F2")
                    else:
                        ws.cell(i, 1).alignment = Alignment(indent=1)
            ws.cell(len(t) + 5, 1, (note + " " if note else "") + "Source : enquête ménage, Zone de Santé de Limete, 2026.").font = Font(italic=True, size=9)
            ws.column_dimensions["A"].width = 58
            for lettre in "BCDEFG":
                ws.column_dimensions[lettre].width = 20
    return buf.getvalue()
